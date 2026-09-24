"""Product import from Excel/CSV: preview (validate) then commit."""

import csv
import io
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from lib.auth import audit, require_admin
from lib.db import db
from models.schemas import ImportCommitIn, ImportResult, ImportRow
from routers.products import generate_unique_barcode, record_movement

router = APIRouter(prefix="/products/import", tags=["import"])

HEADERS = ["SKU", "Barcode", "Nama", "Kategori", "Satuan", "Harga Beli", "Harga Jual", "Stok", "Stok Minimum"]
KEYS = {"sku": "sku", "barcode": "barcode", "nama": "name", "nama produk": "name", "kategori": "category",
        "satuan": "unit", "harga beli": "buy_price", "harga jual": "sell_price", "stok": "stock",
        "stok minimum": "min_stock", "stok min": "min_stock"}


def _num(v) -> float | None:
    if v is None or str(v).strip() == "":
        return 0.0
    try:
        return float(str(v).replace("Rp", "").replace(".", "").replace(",", ".").strip()) if isinstance(v, str) else float(v)
    except ValueError:
        return None


def _read(filename: str, raw: bytes) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = raw.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect)]


@router.get("/template.xlsx")
async def template(_: dict = Depends(require_admin)):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["SBK-010", "", "Tepung Terigu 1 Kg", "Sembako", "pcs", 11000, 13000, 20, 5])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="template-import-produk.xlsx"'})


@router.post("/preview", response_model=list[ImportRow])
async def preview(file: UploadFile = File(...), _: dict = Depends(require_admin)):
    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(400, "File terlalu besar (maks 10 MB)")
    try:
        table = _read(file.filename or "", raw)
    except Exception:
        raise HTTPException(400, "File tidak bisa dibaca. Gunakan .xlsx atau .csv")
    if len(table) < 2:
        raise HTTPException(400, "File kosong")
    cols = {i: KEYS.get(str(h or "").strip().lower()) for i, h in enumerate(table[0])}
    if "sku" not in cols.values() or "name" not in cols.values():
        raise HTTPException(400, "Kolom wajib SKU dan Nama tidak ditemukan di baris pertama")
    existing = {p["sku"]: p async for p in db.products.find({}, {"_id": 0, "sku": 1, "barcode": 1, "id": 1})}
    barcodes = {p["barcode"]: p["sku"] for p in existing.values() if p.get("barcode")}
    seen_sku, seen_bc, out = set(), set(), []
    for n, r in enumerate(table[1:], start=2):
        if not any(str(c or "").strip() for c in r):
            continue
        v = {k: r[i] for i, k in cols.items() if k and i < len(r)}
        errs: list[str] = []
        sku = str(v.get("sku") or "").strip()
        name = str(v.get("name") or "").strip()
        bc = str(v.get("barcode") or "").strip() or None
        if bc and bc.endswith(".0"):
            bc = bc[:-2]
        nums = {k: _num(v.get(k)) for k in ("buy_price", "sell_price", "stock", "min_stock")}
        if not sku:
            errs.append("SKU kosong")
        if not name:
            errs.append("Nama kosong")
        for k, lbl in (("buy_price", "Harga beli"), ("sell_price", "Harga jual"), ("stock", "Stok"), ("min_stock", "Stok minimum")):
            if nums[k] is None or nums[k] < 0:
                errs.append(f"{lbl} tidak valid")
        if sku in seen_sku:
            errs.append("SKU dobel di file")
        if bc and (bc in seen_bc or (bc in barcodes and barcodes[bc] != sku)):
            errs.append("Barcode sudah dipakai")
        seen_sku.add(sku)
        if bc:
            seen_bc.add(bc)
        out.append(ImportRow(row=n, sku=sku, barcode=bc, name=name, category=str(v.get("category") or "").strip(),
                             unit=str(v.get("unit") or "pcs").strip() or "pcs",
                             **{k: (x if x is not None and x >= 0 else 0) for k, x in nums.items()},
                             status="error" if errs else ("update" if sku in existing else "new"), errors=errs))
    return out


@router.post("", response_model=ImportResult)
async def commit(body: ImportCommitIn, admin: dict = Depends(require_admin)):
    cats = {c["name"].lower(): c async for c in db.categories.find({}, {"_id": 0})}
    created = updated = skipped = 0
    new_cats: list[str] = []
    now = datetime.now(timezone.utc)
    for r in body.rows:
        if r.status == "error" or not r.sku or not r.name:
            skipped += 1
            continue
        cat = None
        if r.category:
            cat = cats.get(r.category.lower())
            if not cat:
                cat = {"id": str(uuid.uuid4()), "name": r.category, "description": ""}
                await db.categories.insert_one(dict(cat))
                cats[r.category.lower()] = cat
                new_cats.append(r.category)
        fields = {"name": r.name, "unit": r.unit, "buy_price": r.buy_price, "sell_price": r.sell_price,
                  "min_stock": r.min_stock, "category_id": cat["id"] if cat else None,
                  "category_name": cat["name"] if cat else None, "updated_at": now}
        old = await db.products.find_one({"sku": r.sku}, {"_id": 0})
        if r.barcode and await db.products.find_one({"barcode": r.barcode, "sku": {"$ne": r.sku}}):
            skipped += 1
            continue
        if old:
            if r.barcode:
                fields["barcode"] = r.barcode
            await db.products.update_one({"id": old["id"]}, {"$set": {**fields, "stock": r.stock}})
            if old["stock"] != r.stock:
                await record_movement({**old, **fields}, "adjustment", r.stock - old["stock"], old["stock"], r.stock, admin, "Import produk")
            updated += 1
        else:
            doc = {"id": str(uuid.uuid4()), "sku": r.sku, "barcode": r.barcode or await generate_unique_barcode(),
                   **fields, "stock": r.stock, "supplier": "", "photo_url": "", "active": True, "created_at": now}
            await db.products.insert_one(dict(doc))
            if r.stock:
                await record_movement(doc, "stock_in", r.stock, 0, r.stock, admin, "Import produk (stok awal)")
            created += 1
    await audit(admin, "product_import", f"baru {created}, update {updated}, lewati {skipped}")
    return ImportResult(created=created, updated=updated, skipped=skipped, new_categories=new_cats)
